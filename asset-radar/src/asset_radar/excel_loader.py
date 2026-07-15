from __future__ import annotations

import os
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from .models import RawNotice


def _find_manual_screening_file(base_dir: Path) -> Path:
    for name in os.listdir(base_dir):
        if "2025-12-16" in name and name.endswith(".xlsx"):
            return base_dir / name
    raise FileNotFoundError("Manual screening workbook not found")


def load_marked_notices(base_dir: Path, limit: int | None = None) -> list[RawNotice]:
    path = _find_manual_screening_file(base_dir)
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[1]

    def color_code(color):
        if color is None:
            return None
        if color.type == "rgb":
            return color.rgb
        if color.type == "indexed":
            return f"indexed:{color.indexed}"
        return str(color.type)

    def is_yellow(cell) -> bool:
        codes = {color_code(cell.fill.fgColor), color_code(cell.fill.start_color), color_code(cell.fill.end_color)}
        codes = {code for code in codes if code}
        return any(code in {"FFFFFF00", "FFFFFF99", "FFFFF2CC", "FFFFFFCC", "FFFFEB9C", "indexed:5", "indexed:6", "indexed:13"} for code in codes)

    def is_red_font(cell) -> bool:
        code = color_code(cell.font.color)
        return bool(code in {"FFFF0000", "FFFF3333", "FFCC0000", "indexed:2", "indexed:10"})

    def text(row: int, col: int) -> str:
        value = ws.cell(row, col).value
        return "" if value is None else str(value)

    notices: list[RawNotice] = []
    for row in range(2, ws.max_row + 1):
        marker_cell = ws.cell(row, 8)
        if not (is_yellow(marker_cell) or is_red_font(marker_cell)):
            continue
        title = text(row, 1)
        raw_text = text(row, 8)
        notice = RawNotice(
            source_platform="manual_excel_gold_sample",
            source_url=text(row, 13),
            title=title,
            notice_date=None,
            raw_text=raw_text,
            debtor=text(row, 3),
            disposal_agency=text(row, 9),
            city=text(row, 4) or text(row, 6),
            amount_text=text(row, 2),
            metadata={
                "excel_row": row,
                "yellow": is_yellow(marker_cell),
                "red_font": is_red_font(marker_cell),
                "fit_label": text(row, 7),
                "manual_status": text(row, 14),
                "hit_field": text(row, 16),
                "hit_basis": text(row, 17),
                "loaded_at": date.today().isoformat(),
            },
        )
        notices.append(notice)
        if limit and len(notices) >= limit:
            break
    return notices
