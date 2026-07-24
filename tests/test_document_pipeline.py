from __future__ import annotations

import gzip
import io
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from document_parser import DataPackageParser, choose_parse_route
from storage import LocalStorage, StorageError


class DocumentPipelineTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.storage = LocalStorage(self.root / "storage")

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_local_storage_is_atomic_and_rejects_path_escape(self) -> None:
        self.storage.put_json("cache/result.json", {"status": "ok"})
        self.assertEqual(self.storage.get_json("cache/result.json"), {"status": "ok"})
        self.assertTrue(self.storage.exists("cache/result.json"))
        self.storage.put_file("uploads/large.bin", io.BytesIO(b"chunked-upload"))
        self.assertEqual(self.storage.get_bytes("uploads/large.bin"), b"chunked-upload")
        with self.assertRaises(StorageError):
            self.storage.put_text("../outside.txt", "blocked")

    def test_office_documents_cache_and_element_artifact(self) -> None:
        paths = [
            self._make_docx(),
            self._make_xlsx(),
            self._make_pptx(),
            self._make_markdown(),
        ]
        parser = DataPackageParser(self.storage, document_workers=2, ocr_workers=1)

        first = parser.parse(paths)
        self.assertEqual(first.errors, [])
        self.assertEqual(len(first.documents), len(paths))
        self.assertTrue(all(document.route == "unstructured" for document in first.documents))
        self.assertTrue(all(document.elements for document in first.documents))

        second = parser.parse(paths)
        self.assertEqual(second.errors, [])
        self.assertEqual(sum("命中解析缓存" in line for line in second.logs), len(paths))

        artifact_key = parser.save_elements_artifact("fixture-package", second.documents)
        artifact = self.storage.get_bytes(artifact_key)
        self.assertIsNotNone(artifact)
        rows = [json.loads(line) for line in gzip.decompress(artifact or b"").splitlines()]
        self.assertEqual(len(rows), sum(len(document.elements) for document in second.documents))
        self.assertTrue(all(row["source_file"] for row in rows))

    def test_searchable_pdf_uses_lightweight_parser(self) -> None:
        pdf_path = self._make_searchable_pdf()
        self.assertEqual(choose_parse_route(pdf_path), "text_pdf")

        result = DataPackageParser(self.storage, document_workers=1, ocr_workers=1).parse([pdf_path])
        self.assertEqual(result.errors, [])
        self.assertEqual(result.documents[0].parser, "pypdf")
        self.assertIn("investment", result.documents[0].markdown.lower())

    def test_images_and_scanned_pdf_route_to_baidu_ocr(self) -> None:
        image_path = self.root / "scan.png"
        image_path.write_bytes(b"not-read-during-routing")
        self.assertEqual(choose_parse_route(image_path), "baidu_ocr")

        scanned_pdf = self.root / "scan.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with scanned_pdf.open("wb") as pdf_file:
            writer.write(pdf_file)
        self.assertEqual(choose_parse_route(scanned_pdf), "baidu_ocr")

    def _make_docx(self) -> Path:
        path = self.root / "investment.docx"
        document = Document()
        document.add_heading("Investment Project", level=1)
        document.add_paragraph("Debt principal is 580 million yuan.")
        document.save(path)
        return path

    def _make_xlsx(self) -> Path:
        path = self.root / "finance.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Finance"
        sheet.append(["Item", "Amount"])
        sheet.append(["Debt principal", 580000000])
        workbook.save(path)
        workbook.close()
        return path

    def _make_pptx(self) -> Path:
        path = self.root / "summary.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Project Summary"
        slide.placeholders[1].text = "Asset package and transaction structure"
        presentation.save(path)
        return path

    def _make_markdown(self) -> Path:
        path = self.root / "notes.md"
        path.write_text("# Legal Notes\n\nThe project land is mortgaged.\n", encoding="utf-8")
        return path

    def _make_searchable_pdf(self) -> Path:
        path = self.root / "searchable.pdf"
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = DecodedStreamObject()
        text = "Investment project searchable text with debt principal 580 million and asset details. " * 3
        stream.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
        with path.open("wb") as pdf_file:
            writer.write(pdf_file)
        return path


if __name__ == "__main__":
    unittest.main()
